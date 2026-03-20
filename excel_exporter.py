# modules/excel_exporter.py
#
# Single-output format: Excel (.xlsx) with two sheets.
# Sheet 1: Extraction results (with annotated PDF link)
# Sheet 2: Underlying ISINs + Bloomberg tickers

import logging
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

import config

logger = logging.getLogger(__name__)

# Columns for Sheet 1 -- DENOMINATION removed, ANNOTATED_PDF added
MAIN_COLUMNS = [
    "source_file", "LANGUAGE", "PST_ISIN", "BIL", "CLN",
    "CAPITAL_PROTECTION", "MATURITY", "WORST_OR_AVERAGE",
    "CURRENCY", "ISSUER", "COUPON", "SSPA_TYPE", "ANNOTATED_PDF",
]


def export_excel(records: list, output_path: str) -> Path:
    """
    Export extraction results to an Excel file.

    records: list of dicts, each with keys:
        source_file, language, values (dict), annotated_pdf (str or None)
    """
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()

    # -- Styles --
    header_font = Font(bold=True, color=config.EXCEL_HEADER_FONT_COLOR, size=11)
    header_fill = PatternFill(
        start_color=config.EXCEL_HEADER_COLOR,
        end_color=config.EXCEL_HEADER_COLOR,
        fill_type="solid",
    )
    header_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    # ========== Sheet 1: Results ==========
    ws1 = wb.active
    ws1.title = "Resultats"

    for ci, col_name in enumerate(MAIN_COLUMNS, 1):
        cell = ws1.cell(row=1, column=ci, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    for ri, record in enumerate(records, 2):
        vals = record.get("values", {})
        row_data = {
            "source_file": record.get("source_file", ""),
            "LANGUAGE": record.get("language", ""),
            "ANNOTATED_PDF": record.get("annotated_pdf", ""),
        }
        row_data.update(vals)

        for ci, col_name in enumerate(MAIN_COLUMNS, 1):
            val = row_data.get(col_name, "")
            if val is None:
                val = ""
            if isinstance(val, bool):
                val = "True" if val else "False"
            # Skip list-type values (UNDERLYING_ISINS, BLOOMBERG_TICKERS)
            if isinstance(val, list):
                val = ""
            cell = ws1.cell(row=ri, column=ci, value=str(val))
            cell.border = thin_border

    for ci, col_name in enumerate(MAIN_COLUMNS, 1):
        ws1.column_dimensions[_col_letter(ci)].width = max(len(col_name) + 4, 16)

    # ========== Sheet 2: Underlying ISINs + Tickers ==========
    ws2 = wb.create_sheet("Underlying_ISINs")
    und_cols = ["source_file", "PST_ISIN", "UNDERLYING_ISIN", "BLOOMBERG_TICKER"]

    for ci, col_name in enumerate(und_cols, 1):
        cell = ws2.cell(row=1, column=ci, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    ri = 2
    for record in records:
        vals = record.get("values", {})
        pst = vals.get("PST_ISIN", "")
        underlyings = vals.get("UNDERLYING_ISINS", [])
        tickers = vals.get("BLOOMBERG_TICKERS", [])

        # If we have underlyings, write one row per underlying
        if underlyings:
            for idx, uid in enumerate(underlyings):
                ticker = tickers[idx] if idx < len(tickers) else ""
                ws2.cell(row=ri, column=1, value=record.get("source_file", "")).border = thin_border
                ws2.cell(row=ri, column=2, value=str(pst or "")).border = thin_border
                ws2.cell(row=ri, column=3, value=uid).border = thin_border
                ws2.cell(row=ri, column=4, value=ticker).border = thin_border
                ri += 1
        # If no underlyings but we have tickers, write ticker rows
        elif tickers:
            for ticker in tickers:
                ws2.cell(row=ri, column=1, value=record.get("source_file", "")).border = thin_border
                ws2.cell(row=ri, column=2, value=str(pst or "")).border = thin_border
                ws2.cell(row=ri, column=3, value="").border = thin_border
                ws2.cell(row=ri, column=4, value=ticker).border = thin_border
                ri += 1

    for ci in range(1, len(und_cols) + 1):
        ws2.column_dimensions[_col_letter(ci)].width = 30

    wb.save(str(path))
    logger.info(f"Excel exported: {path}")
    return path


def _col_letter(ci: int) -> str:
    """Convert 1-based column index to Excel letter."""
    result = ""
    while ci > 0:
        ci, remainder = divmod(ci - 1, 26)
        result = chr(65 + remainder) + result
    return result
